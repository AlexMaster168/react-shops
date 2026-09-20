"""Снимает скриншоты запущенного приложения (vite preview) и собирает docs/report.docx.

Запуск: pnpm build && pnpm preview (в другом терминале), затем python scripts/make_report.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Cm
from playwright.sync_api import sync_playwright

URL = "http://localhost:4173/"
OUT = Path("docs")
SHOTS = OUT / "screens"
SHOTS.mkdir(parents=True, exist_ok=True)

shots: list[tuple[str, str]] = []


def snap(page, name: str, caption: str, full: bool = True):
    path = SHOTS / f"{name}.png"
    page.screenshot(path=str(path), full_page=full)
    shots.append((str(path), caption))


with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page(viewport={"width": 1280, "height": 900})
    page.goto(URL)
    page.wait_for_selector(".card")

    snap(page, "01-catalog", "Каталог: первая страница (8 товаров), обложки загружаются локально")

    # корзина: штучное добавление
    cards = page.locator(".card")
    cards.nth(0).get_by_role("button", name="В корзину").click()
    cards.nth(0).get_by_label("Добавить ещё одну").click()
    cards.nth(1).get_by_role("button", name="В корзину").click()
    snap(page, "02-card-qty", "Карточка товара: кнопки «−» / «+» и счётчик штук")

    page.get_by_text("Корзина (").click()
    page.wait_for_selector(".popup")
    snap(page, "03-cart", "Корзина: количество по позициям, +/−, удаление, итоговая сумма", full=False)
    page.get_by_label("Убрать одну").last.click()
    snap(page, "04-cart-minus", "Корзина: после нажатия «−» у второй позиции (было 1 шт. — позиция удалена)", full=False)
    page.mouse.click(10, 10)

    # сортировка
    page.locator(".dropdown", has_text="По умолчанию").click()
    page.get_by_role("option", name="Цена: сначала дорогие").click()
    snap(page, "05-sort", "Сортировка: цена по убыванию")

    # пагинация
    page.get_by_role("link", name="2").or_(page.get_by_text("2", exact=True)).last.click()
    snap(page, "06-page2", "Пагинация: вторая страница")

    # фильтры
    page.get_by_role("button", name="Сбросить фильтры").click()
    page.get_by_placeholder("Название или автор...").fill("роберт")
    snap(page, "07-search", "Поиск по названию/автору: «роберт»")

    page.get_by_role("button", name="Сбросить фильтры").click()
    page.get_by_placeholder("до ", exact=False).fill("150")
    page.locator(".rating .icon").nth(3).click()  # рейтинг от 4
    snap(page, "08-filters", "Фильтры: цена до 150 грн и рейтинг от 4")

    page.get_by_placeholder("Название или автор...").fill("нет такой книги")
    snap(page, "09-empty", "Пустая выдача: сообщение «Ничего не найдено»")
    browser.close()

doc = Document()
doc.add_heading("Магазин книг — отчёт со скриншотами", 0)
doc.add_paragraph("React 18 + TypeScript + Redux Toolkit + Vite, пакетный менеджер pnpm.")
for path, caption in shots:
    doc.add_heading(caption, level=2)
    doc.add_picture(path, width=Cm(16))
doc.save(OUT / "report.docx")
print("saved", OUT / "report.docx", len(shots), "screens")
